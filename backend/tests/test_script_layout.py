import io

from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill

from app.modules.workflow.script_layout import build_layout
from app.core.database import connection
from test_numbered_multimodal_script import setup, workbook, upload
from test_script_versions import K


def test_layout_retains_merges_styles_geometry_drawings_and_formula_limits(client,monkeypatch):
    task=setup(client)
    wb=load_workbook(io.BytesIO(workbook()));ws=wb.active
    ws.merge_cells('A1:B1');ws['A1'].font=Font(name='宋体',bold=True,color='FF123456',size=14)
    ws['A1'].fill=PatternFill('solid',fgColor='FFFFCC00');ws.column_dimensions['A'].width=25;ws.row_dimensions[1].height=36
    ws['C2']='=HYPERLINK("https://untrusted.example/", "不要执行")'
    ws.column_dimensions['C'].hidden=True
    out=io.BytesIO();wb.save(out);wb.close()
    file_id=upload(client,out.getvalue())
    saved=client.post(f'/api/v1/tasks/{task}/scripts',json={'inputType':'file','fileId':file_id},headers={'Idempotency-Key':K()}).json()['data']
    path=f"/api/v1/tasks/{task}/scripts/{saved['scriptVersionId']}/layout"
    response=client.get(path);assert response.status_code==200,response.text
    data=response.json()['data'];sheet=data['sheets'][0]
    assert sheet['name']=='脚本' and sheet['merges']==[{'r1':1,'c1':1,'r2':1,'c2':2}]
    assert sheet['rows'][0]==48 and sheet['columns'][0]==180
    title=next(c for c in sheet['cells'] if c['address']=='A1')
    assert title['text']=='品牌A产品展示' and title['style']['fontWeight']==700
    assert title['style']['backgroundColor']=='#ffcc00' and title['style']['color']=='#123456'
    formula=next(c for c in sheet['cells'] if c['address']=='C2')
    assert formula['uncached'] and formula['text']=='' and 'HYPERLINK' in formula['formula']
    assert sheet['hiddenColumns']==[3]
    assert sheet['images'][0]['cell']=='D5' and sheet['images'][0]['width']==100
    assert sheet['images'][0]['height']==60
    image_id=sheet['images'][0]['id']
    assert client.get(path+'/images/'+image_id).headers['content-type']=='image/png'
    assert client.get(path+'/images/invented-image').status_code==404
    with connection(client.app.state.settings.database_path) as db:
        assert db.execute('SELECT COUNT(*) FROM workflow_script_reviews WHERE task_id=?',(task,)).fetchone()[0]==0
    # Another valid task may not read or request images from this immutable version.
    with connection(client.app.state.settings.database_path) as db:p=db.execute('SELECT project_id FROM creator_tasks WHERE id=?',(task,)).fetchone()[0]
    other=client.post(f'/api/v1/projects/{p}/creators',json={'creatorName':'另一个达人','platform':'douyin'},headers={'Idempotency-Key':K()}).json()['data']['id']
    assert client.get(path.replace(task,other)).status_code==404
    assert client.get((path+'/images/'+image_id).replace(task,other)).status_code==404


def test_layout_limits_are_explicit_and_no_formula_is_recalculated():
    wb=Workbook();ws=wb.active;ws['A1']='=1+1';ws['AG501']='超出范围'
    extra=wb.create_sheet('另一个大表');extra['AF500']='不静默忽略'
    out=io.BytesIO();wb.save(out);wb.close()
    data=build_layout(out.getvalue(),{'images':[],'uncovered':[]})
    assert len(data['sheets'])==1
    assert data['sheets'][0]['rowCount']==500 and data['sheets'][0]['colCount']==32
    assert any('仅预览' in w for w in data['warnings'])
    assert any('网格上限' in w for w in data['warnings'])
    assert next(c for c in data['sheets'][0]['cells'] if c['address']=='A1')['text']==''
