import io
import zipfile

from docx import Document
from docx.oxml.ns import qn

from app.modules.workflow.word_preview import add_source_bookmarks, source_bookmark_name


def test_bookmarks_follow_parser_coordinates_and_merged_cell_aliases():
    document=Document()
    document.add_paragraph('重复文本')
    table=document.add_table(rows=3,cols=3)
    table.cell(0,0).merge(table.cell(0,1)).text='横向合并'
    table.cell(1,0).merge(table.cell(2,0)).text='纵向合并'
    table.cell(2,2).text='重复文本'
    document.add_paragraph('表后段落')
    raw=io.BytesIO();document.save(raw)
    result,count=add_source_bookmarks(raw.getvalue(),'12345678abcdef')
    assert count==11
    marked=Document(io.BytesIO(result))
    def names(element):
        return {node.get(qn('w:name')) for node in element.iter(qn('w:bookmarkStart'))}
    name=lambda location:source_bookmark_name('12345678abcdef',location)
    assert name('paragraph:0') in names(marked.paragraphs[0]._p)
    assert name('paragraph:2') in names(marked.paragraphs[1]._p)
    assert {name('table:0:r0c0'),name('table:0:r0c1')}<=names(marked.tables[0].cell(0,0)._tc)
    assert {name('table:0:r1c0'),name('table:0:r2c0')}<=names(marked.tables[0].cell(1,0)._tc)
    assert name('table:0:r2c2') in names(marked.tables[0].cell(2,2)._tc)
    ids=[node.get(qn('w:id')) for node in marked.element.iter(qn('w:bookmarkStart'))]
    assert len(ids)==len(set(ids))
    assert add_source_bookmarks(result,'12345678abcdef')==(result,0)
    with zipfile.ZipFile(io.BytesIO(raw.getvalue())) as original,zipfile.ZipFile(io.BytesIO(result)) as preview:
        assert original.namelist()==preview.namelist()
        for member in original.namelist():
            if member!='word/document.xml':assert original.read(member)==preview.read(member)
