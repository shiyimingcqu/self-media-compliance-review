import copy
import io
import json

import httpx
import pytest
from openpyxl import Workbook
from openpyxl.drawing.image import Image as SheetImage
from PIL import Image

from app.modules.workflow import router as workflow
from app.modules.workflow.prompts import prompt
from app.modules.workflow.script_image_context import annotate_image_context
from test_numbered_multimodal_script import setup, upload, report, submit


def test_word_reference_column_proves_purpose_and_links_actual_plan_row():
    blocks=[{'paragraphId':'header','location':'table:1:r1c4','text':'参考图'},
            {'paragraphId':'plan','location':'table:1:r2c2','text':'本次拿推广咖啡，只参考偷笑动作'},
            {'paragraphId':'notes','location':'table:1:r2c3','text':'侧重反转感'}]
    manifest={'images':[{'id':'image','ref':'i1','locations':['table:1:r2c4']}],'uncovered':[]}
    original=copy.deepcopy(manifest)
    output=annotate_image_context(manifest,blocks)
    image=output['images'][0]
    assert image['purpose']=='reference'
    assert image['contexts'][0]=={'location':'table:1:r2c4','purpose':'reference','label':'参考图',
                                  'labelLocation':'table:1:r1c4','labelBlockId':'header','rowBlockIds':['plan','notes']}
    assert manifest==original


def test_reference_header_does_not_label_unrelated_tables_columns_or_unknown_positions():
    blocks=[{'paragraphId':'p1','location':'table:1:r1c4','text':'参考图'}]
    manifest={'images':[{'id':str(i),'locations':[location]} for i,location in enumerate(
        ['table:2:r2c4','table:1:r2c2','位置未知'])]}
    assert all(image['purpose']=='unknown' for image in annotate_image_context(manifest,blocks)['images'])


def test_mixed_use_and_closest_explicit_header_are_preserved():
    blocks=[{'paragraphId':'p1','location':'脚本!D1','text':'参考图'},
            {'paragraphId':'p2','location':'脚本!D4','text':'本次拍摄画面'}]
    manifest={'images':[{'id':'a','locations':['脚本!D2']},{'id':'b','locations':['脚本!D5']},
                        {'id':'c','locations':['脚本!D2','脚本!D5']}]}
    images=annotate_image_context(manifest,blocks)['images']
    assert [image['purpose'] for image in images]==['reference','planned','mixed']
    assert images[1]['contexts'][0]['labelLocation']=='脚本!D4'


def test_prompt_distinguishes_reference_props_from_adopted_content_and_rule_scope():
    text=prompt('script-review')['content']
    for phrase in ('【先区分参考素材与本次拍摄安排】','不能要求参考图也必须展示推广品牌',
                   '也不得仅以“如果实拍沿用就违规”判待确认','不等于“所有非麦当劳食品”',
                   '本次仍拿竞品奶茶并展示标识','不能把所有参考图片一概跳过',
                   '不是自动豁免','返回compliant'):
        assert phrase in text
    for key in ('brief-script-standards','brief-script-rules-text'):
        generated=prompt(key)['content']
        assert '禁止其他饮料不能扩写为禁止所有食品' in generated
        assert '【先区分参考素材与本次拍摄安排】' not in generated  # 不向生成接口注入审稿输出字段


def reference_workbook(adopted):
    image=io.BytesIO();Image.new('RGB',(60,60),'red').save(image,format='PNG');image.seek(0)
    wb=Workbook();ws=wb.active;ws.title='脚本'
    ws['A1']='画面';ws['D1']='参考图'
    ws['A2']='本次仍拿竞品奶茶并展示标识' if adopted else '本次拿推广咖啡，只借用参考特效，道具换成指定产品'
    ws.add_image(SheetImage(image),'D2')
    output=io.BytesIO();wb.save(output);return output.getvalue()


@pytest.mark.parametrize('adopted,verdict',[(False,'compliant'),(True,'non_compliant')])
def test_context_reaches_actual_provider_and_is_frozen_without_auto_overriding_judgment(client,monkeypatch,adopted,verdict):
    task=setup(client);file_id=upload(client,reference_workbook(adopted));actual=httpx.Client;calls=[]
    def provider(request):
        body=json.loads(request.content);calls.append(body)
        material=json.loads(body['messages'][1]['content'][0]['text'])
        image=material['multimodal']['images'][0]
        assert image['purpose']=='reference' and image['contexts'][0]['label']=='参考图'
        assert image['contexts'][0]['rowBlockIds']
        assert '不能要求参考图也必须展示推广品牌' in body['messages'][0]['content']
        assert '【创意表达的实务容忍口径】' in body['messages'][0]['content']
        assert '逐字引文只证明说了什么' in body['messages'][0]['content']
        value=report(material)
        value['findings'][1].update(verdict=verdict,image_ids=[],description='依据本次画面安排判断，不把参考图原道具当作拟拍道具。')
        value['image_checks'][0].update(rule_ids=[],description='该图用于特效风格参考，不证明本次道具或品牌。')
        return httpx.Response(200,json={'choices':[{'finish_reason':'stop','message':{'content':json.dumps(value)}}]})
    monkeypatch.setattr(workflow.httpx,'Client',lambda **kw:actual(transport=httpx.MockTransport(provider),**kw))
    response=submit(client,task,file_id)
    assert response.status_code==201,response.text
    review=response.json()['data']['autoReview']
    assert review['status']=='succeeded',review
    assert review['snapshot']['multimodal']['images'][0]['purpose']=='reference'
    assert review['result']['report']['findings'][1]['verdict']==verdict
    assert review['acceptance']['passed'] is (not adopted)
    assert len(calls)==1
