"""t-38 真实格式解析（XLSX/DOCX）单元测试。

覆盖 A 13:01 两处修正：
1) location 精确到真实单元格/行列（口播 D12 必须仍指 D12；表格至少 table/row/column）；
2) 未识读图片按"媒体 + 引用位置"去重（WPS 同一图不重复计两张），并置 coverage=partial；
   历史缺元数据 → unknown（不补 false）。
压缩包安全门另测。CI 用 openpyxl/python-docx 现场生成文档，不依赖用户私有样本。
"""

import io
import zipfile

import pytest

from app.modules.brief_script.parsing.doc import COVERAGE_FULL, COVERAGE_PARTIAL, DocFormatError, open_zip_guarded
from app.modules.brief_script.parsing.docx import parse_docx
from app.modules.brief_script.parsing.xlsx import parse_xlsx


def _xlsx_bytes(rows, merges=(), with_image=False):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "视频脚本"
    for row in rows:
        ws.append(row)
    for m in merges:
        ws.merge_cells(m)
    buf = io.BytesIO()
    wb.save(buf)
    data = buf.getvalue()
    if with_image:
        src = zipfile.ZipFile(io.BytesIO(data))
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
            for name in src.namelist():
                dst.writestr(name, src.read(name))
            dst.writestr("xl/media/image1.png", b"\x89PNG\r\n\x1a\n" + b"\x00" * 32)
        data = out.getvalue()
    return data


def _docx_bytes(paragraphs, table=None, with_media=False):
    from docx import Document
    d = Document()
    for p in paragraphs:
        d.add_paragraph(p)
    if table:
        t = d.add_table(rows=len(table), cols=len(table[0]))
        for r_i, row in enumerate(table):
            for c_i, val in enumerate(row):
                t.cell(r_i, c_i).text = val
    buf = io.BytesIO()
    d.save(buf)
    data = buf.getvalue()
    if with_media:
        src = zipfile.ZipFile(io.BytesIO(data))
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
            for name in src.namelist():
                dst.writestr(name, src.read(name))
            dst.writestr("word/media/image1.png", b"\x89PNG\r\n\x1a\n" + b"\x00" * 32)
        data = out.getvalue()
    return data


def test_xlsx_each_nonempty_cell_has_exact_coordinate():
    """A 13:01 修正1：每非空 cell 一个块，location 精确到真实单元格（不拿 A 列冒充整行）。"""
    data = _xlsx_bytes([["时长", "场景", "画面内容", "台词/画外音", "画面参考"]])
    doc = parse_xlsx(data)
    locs = [b.location for b in doc.blocks]
    assert locs == ["视频脚本!A1", "视频脚本!B1", "视频脚本!C1", "视频脚本!D1", "视频脚本!E1"]
    by_loc = {b.location: b.text for b in doc.blocks}
    assert by_loc["视频脚本!D1"] == "台词/画外音"   # 口播列就是 D，不是 A
    assert by_loc["视频脚本!E1"] == "画面参考"


def test_xlsx_data_row_keeps_own_cell_identity():
    """数据行：D12 口播文本必须仍标 视频脚本!D12。"""
    rows = [["头"]] + [[""] for _ in range(10)] + [["", "", "", "妈妈：冰奶铁好喝啊", ""]]
    doc = parse_xlsx(_xlsx_bytes(rows))
    hit = [b for b in doc.blocks if "冰奶铁" in b.text]
    assert hit and hit[0].location == "视频脚本!D12"


def test_xlsx_merge_anchor_keeps_own_coordinate_and_warns():
    data = _xlsx_bytes([["主页", "x"], ["时长", "场景", "画面内容"]], merges=["A1:B1"])
    doc = parse_xlsx(data)
    assert any("合并区" in w for w in doc.warnings)
    texts = {b.location: b.text for b in doc.blocks}
    assert texts["视频脚本!A2"] == "时长" and texts["视频脚本!B2"] == "场景"  # 表头各自独立


def test_xlsx_same_image_not_double_counted():
    """A 13:01 修正2：WPS 同一图（media + cellimages 引用）只记一处，带引用锚点。"""
    src = zipfile.ZipFile(io.BytesIO(_xlsx_bytes([["a"]], with_image=True)))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for name in src.namelist():
            dst.writestr(name, src.read(name))
        dst.writestr("xl/cellimages.xml",
                     '<etc:cellImages xmlns:etc="http://www.wps.cn/officeDocument/2017/etCustomData" '
                     'xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" '
                     'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
                     'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
                     '<etc:cellImage><xdr:pic><xdr:nvPicPr><xdr:cNvPr id="2" name="ID_ABC"/></xdr:nvPicPr>'
                     '<xdr:blipFill><a:blip r:embed="rId1"/></xdr:blipFill></xdr:pic></etc:cellImage></etc:cellImages>')
        dst.writestr("xl/_rels/cellimages.xml.rels",
                     '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                     '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
                     'relationships/image" Target="media/image1.png"/></Relationships>')
    doc = parse_xlsx(out.getvalue())
    assert doc.coverage == COVERAGE_PARTIAL and doc.partial is True
    assert len(doc.uncovered_media) == 1, doc.uncovered_media          # 一张图只一次
    assert "image1.png" in doc.uncovered_media[0]
    assert "ID_ABC" in doc.uncovered_media[0]                          # 保留引用映射位置


def test_xlsx_full_text_coverage_when_no_media():
    doc = parse_xlsx(_xlsx_bytes([["纯文本"]]))
    assert doc.coverage == COVERAGE_FULL and doc.uncovered_media == []


def test_docx_table_cell_has_table_row_column():
    """A 13:01 修正1：Word 表格至少 table/row/column 定位。"""
    doc = parse_docx(_docx_bytes(["大纲"], table=[["时长", "特效"], ["6S", "偷笑女"]]))
    locs = [b.location for b in doc.blocks]
    assert "table:0:r0c0" in locs and "table:0:r0c1" in locs
    assert "table:0:r1c1" in locs
    by_loc = {b.location: b.text for b in doc.blocks}
    assert by_loc["table:0:r1c1"] == "偷笑女"


def test_docx_paragraph_location_and_media_partial():
    doc = parse_docx(_docx_bytes(["正文一", "正文二"], with_media=True))
    locs = [b.location for b in doc.blocks]
    assert "paragraph:0" in locs and "paragraph:1" in locs
    assert doc.coverage == COVERAGE_PARTIAL and doc.uncovered_media


def test_zip_guard_rejects_traversal_and_member_bomb():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("../evil.xml", "<a/>")
    with pytest.raises(DocFormatError):
        open_zip_guarded(buf.getvalue())

    buf2 = io.BytesIO()
    with zipfile.ZipFile(buf2, "w") as z:
        for i in range(2100):
            z.writestr(f"f{i}.xml", "<a/>")
    with pytest.raises(DocFormatError):
        open_zip_guarded(buf2.getvalue())


def test_non_zip_payload_rejected():
    with pytest.raises(DocFormatError):
        open_zip_guarded(b"not a zip at all")


def test_xlsx_formula_without_cache_is_partial_not_silently_dropped():
    """A 13:53①：data_only=True 取缓存值，缺缓存会读成 None；不能因看不到 '=' 就静默丢。"""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "表单"
    ws["A1"] = "标题文本"
    ws["B1"] = "=A1&\"-后缀\""     # 写入公式但没有保存过计算缓存
    buf = io.BytesIO()
    wb.save(buf)
    doc = parse_xlsx(buf.getvalue())
    assert doc.coverage == COVERAGE_PARTIAL          # 不因缺缓存而报 full
    assert any("公式" in w for w in doc.warnings)
    assert any("B1" in r for r in doc.uncovered_other)  # 精确定位到 B1
    assert "标题文本" in [b.text for b in doc.blocks]    # 已知文本仍保留


def test_partial_reason_survives_finalize_without_media():
    """A 13:53②：finalize 不得把非图片类 partial 原因抹掉（原实现 partial=bool(media)）。"""
    from app.modules.brief_script.parsing.doc import ParsedDoc, finalize_media, mark_partial
    doc = ParsedDoc(blocks=[])
    mark_partial(doc, "超限截断")
    finalize_media(doc)                              # 无 media
    assert doc.partial is True
    assert doc.coverage == COVERAGE_PARTIAL
    assert doc.uncovered_other == ["超限截断"]


def test_docx_preserves_paragraph_table_paragraph_order():
    """A 13:53③：必须按 body 原顺序（段落-表格-段落），不能先全段落再全表格重排。"""
    from docx import Document
    d = Document()
    d.add_paragraph("第一段")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "表头A"
    t.cell(0, 1).text = "表头B"
    d.add_paragraph("第三段")
    buf = io.BytesIO()
    d.save(buf)
    doc = parse_docx(buf.getvalue())
    locs = [b.location for b in doc.blocks]
    assert locs == ["paragraph:0", "table:0:r0c0", "table:0:r0c1", "paragraph:2"], locs


def test_docx_paragraph_overflow_is_partial_not_full():
    """A 13:53②：段落超上限不得静默截断后报 full。"""
    import app.modules.brief_script.parsing.docx as m
    data = _docx_bytes([f"段{i}" for i in range(6)])
    old = m.MAX_PARAGRAPHS
    m.MAX_PARAGRAPHS = 3
    try:
        doc = m.parse_docx(data)
    finally:
        m.MAX_PARAGRAPHS = old
    assert doc.coverage == COVERAGE_PARTIAL
    assert any("上限" in r for r in doc.uncovered_other)
    assert doc.uncovered_other
