import base64
import io
import json
import zipfile

import httpx
import pytest
from openpyxl import Workbook
from openpyxl.drawing.image import Image as SheetImage
from PIL import Image

from app.modules.workflow import router as workflow
from app.modules.workflow import script_media
from app.modules.brief_script.parsing.xlsx import parse_xlsx
from test_separate_standards import project, payload
from test_script_versions import K


def workbook(image_only=False):
    image = io.BytesIO(); Image.new('RGB', (100, 60), 'red').save(image, format='PNG'); image.seek(0)
    wb = Workbook(); ws = wb.active; ws.title = '脚本'
    if not image_only:
        ws['A1'] = '品牌A产品展示'; ws['B5'] = '镜头参考'
    ws.add_image(SheetImage(image), 'D5')
    data = io.BytesIO(); wb.save(data)
    return data.getvalue()


def setup(client):
    p = project(client); body = payload(client)
    body['scriptRules'] = [
        {'text':'品牌信息正确','criterion':'文字与图片使用品牌A','category':'品牌','importance':'must'},
        {'text':'图片不得出现竞品','criterion':'检查参考图片是否出现竞品包装','category':'图片','importance':'must'}]
    r = client.post(f"/api/v1/projects/{p['id']}/publish-task", json=body, headers={'Idempotency-Key':K()})
    assert r.status_code == 201, r.text
    task = r.json()['data']['id']
    client.put('/api/v1/settings/models', json={'endpoint':'https://models.example.com/v1','model':'omni-test','revision':0,'apiKey':'sk-test-only'})
    return task


def upload(client, data):
    r=client.post('/api/v1/files', files={'file':('script.xlsx',data,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}, data={'purpose':'script'},headers={'Idempotency-Key':K()})
    assert r.status_code == 201, r.text
    return r.json()['data']['id']


def report(snapshot, non_compliant=False):
    rules = snapshot['rules']; images = snapshot['multimodal']['images']
    # 模型只看得见短引用 ref（长哈希抄写易错），服务端负责还原真实图片 id；
    # 历史快照没有 ref 时仍回落到真实 id，两种写法都接受。
    return {'summary':'已结合文字与图片逐条检查','findings':[
        {'rule_id':r['id'],'rule_number':i,'verdict':'non_compliant' if non_compliant and i==2 else 'compliant',
         'image_ids':[im.get('ref') or im['id'] for im in images],'quote':'','description':'第2条图片存在竞品' if non_compliant and i==2 else '已核对本条', 'suggestion':'替换图片' if non_compliant and i==2 else ''}
        for i,r in enumerate(rules,1)],
        'image_checks':[{'image_id':im.get('ref') or im['id'],'description':'脚本!D5 参考图片为红色画面，已按两条规则核对','rule_ids':[r['id'] for r in rules]} for im in images]}


def submit(client, task, file_id, key=None):
    return client.post(f'/api/v1/tasks/{task}/scripts/submit-and-review',json={'inputType':'file','fileId':file_id,'configRevision':1,'allowRealCall':True},headers={'Idempotency-Key':key or K()})


def test_numbered_rules_and_excel_images_really_in_multimodal_request(client, monkeypatch):
    task=setup(client); file_id=upload(client,workbook()); calls=[]
    actual_client=httpx.Client
    def provider(request):
        calls.append(1); body=json.loads(request.content)
        assert body['model']=='omni-test'
        assert body['reasoning_effort']=='none'
        # 判定类调用必须显式钉温度，否则供应商默认采样会导致同一输入多次判定漂移。
        assert body['temperature']==0
        assert request.extensions['timeout']['read']==300
        parts=body['messages'][1]['content']; material=json.loads(parts[0]['text'])
        assert [r['ruleNumber'] for r in material['rules']]==[1,2]
        assert material['multimodal']['images'][0]['locations']==['脚本!D5']
        assert any(b.get('location')=='脚本!B5' for b in material['scriptBlocks'])
        url=next(p['image_url']['url'] for p in parts if p['type']=='image_url')
        with Image.open(io.BytesIO(base64.b64decode(url.split(',')[1]))) as image:
            assert image.size==(100,60) and image.getpixel((0,0))==(255,0,0)
        return httpx.Response(200,json={'id':'mock-omni','usage':{'total_tokens':123},'choices':[{'message':{'content':json.dumps(report(material,True))}}]})
    monkeypatch.setattr(workflow.httpx,'Client',lambda **kw:actual_client(transport=httpx.MockTransport(provider),**kw))
    key=K(); response=submit(client,task,file_id,key)
    assert response.status_code==201,response.text
    review=response.json()['data']['autoReview']
    assert review['status']=='succeeded',review
    assert review['result']['report']['findings'][1]['rule_number']==2
    assert not review['acceptance']['passed'] and review['acceptance']['checks'][1]['rule_number']==2
    assert 'base64' not in json.dumps(review['snapshot'])
    assert client.get('/api/v1/workflow/review-queue').json()['data']==[]
    image_id=review['snapshot']['multimodal']['images'][0]['id']
    preview=client.get(f"/api/v1/workflow/script-reviews/{review['id']}/images/{image_id}")
    assert preview.status_code==200 and preview.headers['content-type']=='image/png'
    assert client.get(f"/api/v1/workflow/script-reviews/{review['id']}/images/unknown").status_code==404
    assert submit(client,task,file_id,key).json()['data']['autoReview']['id']==review['id']
    assert len(calls)==1


@pytest.mark.parametrize('failure,expected', [('read','PROVIDER_READ_TIMEOUT'),('connect','PROVIDER_CONNECT_TIMEOUT'),('auth','PROVIDER_HTTP_401')])
def test_script_provider_failures_are_diagnosed_without_retry(client, monkeypatch, failure, expected):
    task=setup(client); file_id=upload(client,workbook()); calls=[]
    actual_client=httpx.Client
    def provider(request):
        calls.append(1)
        if failure=='read':raise httpx.ReadTimeout('mock',request=request)
        if failure=='connect':raise httpx.ConnectTimeout('mock',request=request)
        return httpx.Response(401,headers={'x-request-id':'mock-auth-id'},json={'error':{'message':'Key sk-test-only rejected'}})
    monkeypatch.setattr(workflow.httpx,'Client',lambda **kw:actual_client(transport=httpx.MockTransport(provider),**kw))
    key=K(); review=submit(client,task,file_id,key).json()['data']['autoReview']
    assert review['status']=='failed' and review['errorCode']==expected
    assert not review['acceptance']['passed']
    assert 'sk-test-only' not in json.dumps(review)
    assert submit(client,task,file_id,key).json()['data']['autoReview']['id']==review['id']
    assert len(calls)==1
    if failure=='auth': assert review['result']['providerRequestIds']==['mock-auth-id']
    else: assert review['result']['diagnostics']['errorType'].endswith('Timeout')


def test_non_contiguous_quotes_are_preserved_as_verified_segments():
    text='第一段真实台词\n这里是被省略的中间内容\n第二段真实台词'
    finding=workflow.ScriptFinding(rule_id='r',verdict='non_compliant',description='问题',quote='第一段真实台词\n第二段真实台词')
    workflow.validate_quote(finding,text)
    assert finding.quote=='第一段真实台词'
    assert finding.quote_segments==['第一段真实台词','第二段真实台词']
    for fake in ('第一段真实台词\n伪造台词','第二段真实台词\n第一段真实台词','第一段真实台词改写'):
        with pytest.raises(ValueError):
            workflow.validate_quote(workflow.ScriptFinding(rule_id='r',verdict='compliant',description='问题',quote=fake),text)


def test_cached_response_can_be_revalidated_without_provider_call(client,monkeypatch):
    from app.core.database import connection
    task=setup(client);file_id=upload(client,workbook())
    def model(c,k,s):
        value=report(s)
        return {'report':value,'modelResponse':json.dumps(value)}
    monkeypatch.setattr(workflow,'call_script_model',model)
    review=submit(client,task,file_id).json()['data']['autoReview']
    assert review['status']=='succeeded'
    with connection(client.app.state.settings.database_path) as db:
        value=json.loads(db.execute('SELECT result_json FROM workflow_script_reviews WHERE id=?',(review['id'],)).fetchone()[0]);value.pop('report')
        raw=json.loads(value['modelResponse'])
        raw['findings'][0]['quote']='这是模型改写而非原文'
        value['modelResponse']=json.dumps(raw)
        db.execute("UPDATE workflow_script_reviews SET status='failed',error_code='SCRIPT_REPORT_INVALID',result_json=? WHERE id=?",(json.dumps(value),review['id']))
    def forbidden(*args):raise AssertionError('No paid request during revalidation')
    monkeypatch.setattr(workflow,'call_script_model',forbidden)
    restored=workflow.revalidate_saved_script_report(client.app.state.settings,review['id'])
    assert restored['status']=='succeeded' and restored['snapshot']==review['snapshot']
    assert restored['result']['revalidation']['providerResubmitted'] is False
    assert restored['result']['report']['findings'][0]['evidence_status']=='unverified'
    # 服务端把模型返回的短引用还原成快照里的真实图片 id，顺序与内容一一对应
    assert [c['image_id'] for c in restored['result']['report']['image_checks']] == [im['id'] for im in review['snapshot']['multimodal']['images']]
    assert [c['description'] for c in restored['result']['report']['image_checks']] == [c['description'] for c in raw['image_checks']]
    assert not restored['acceptance']['passed']
    assert restored['result']['modelResponse']==json.dumps(raw)
    with connection(client.app.state.settings.database_path) as db:
        assert db.execute("SELECT COUNT(*) FROM audit_events WHERE entity_id=? AND action='saved_response_revalidated'",(review['id'],)).fetchone()[0]==1


def test_exact_full_quote_and_contained_highlights_are_consistent():
    text='女儿：给你买的冰奶铁\n妈妈：喝一口\n妈妈：好喝啊！'
    finding=workflow.ScriptFinding(rule_id='r',verdict='non_compliant',description='问题',quote=text,quote_segments=['女儿：给你买的冰奶铁','妈妈：好喝啊！'])
    workflow.validate_quote(finding,text)
    assert finding.quote=='女儿：给你买的冰奶铁'
    assert finding.quote_segments==['女儿：给你买的冰奶铁','妈妈：好喝啊！']
    for quote,segments in [
        (text,['妈妈：好喝啊！','女儿：给你买的冰奶铁']),
        (text,['伪造台词']),
        ('妈妈：喝一口',['女儿：给你买的冰奶铁']),
        ('女儿：给你买的冰奶铁',['妈妈：喝一口','妈妈：好喝啊！']),
    ]:
        with pytest.raises(ValueError):workflow.validate_quote(workflow.ScriptFinding(rule_id='r',verdict='compliant',description='问题',quote=quote,quote_segments=segments),text)


@pytest.mark.parametrize('bad_section', ['findings', 'mandatory_checks'])
@pytest.mark.parametrize('paragraph_mode', [False, True])
def test_bad_quote_isolated_and_cannot_grant_mandatory_acceptance(bad_section,paragraph_mode):
    snapshot={'rules':[{'id':'r1','ruleNumber':1,'importance':'must','text':'产品名称'},
                       {'id':'r2','ruleNumber':2,'importance':'must','text':'品牌露出'}],
              'script':'妈妈：诶诶诶！产品名称正确','paragraphMode':paragraph_mode,
              'multimodal':{'images':[{'id':'im'}]}}
    value=report(snapshot)
    value['mandatory_checks']=[{'rule_id':r['id'],'requirement':r['text'],'verdict':'compliant','reason':'已核对'} for r in snapshot['rules']]
    value[bad_section][0].update(quote='妈妈：诶诶！产品名称正确',quote_segments=['妈妈：诶诶！产品名称正确'])
    value['findings'][1]['quote']='产品名称正确'
    checked=workflow.validate_script_report(value,snapshot)
    for section in ('findings','mandatory_checks'):
        assert checked[section][0]['verdict']=='needs_confirmation'
        assert checked[section][0]['evidence_status']=='unverified'
        assert checked[section][0]['unverified_quote']=='妈妈：诶诶！产品名称正确'
    assert checked[bad_section][0]['quote']=='' and checked[bad_section][0]['quote_segments']==[]
    assert checked['findings'][1]['quote']=='产品名称正确'
    assert checked['findings'][1]['verdict']=='compliant'
    assert checked['image_checks']==value['image_checks']
    acceptance=workflow.script_acceptance({'status':'succeeded','snapshot_json':json.dumps(snapshot),'result_json':json.dumps({'report':checked})})
    assert not acceptance['passed'] and acceptance['blockingCount']==1


def test_wps_cell_images_and_dispimg_are_mapped_without_executing_formula():
    data=workbook()
    with zipfile.ZipFile(io.BytesIO(data)) as original:
        out=io.BytesIO()
        with zipfile.ZipFile(out,'w') as z:
            for name in original.namelist():
                content=original.read(name)
                if name=='xl/worksheets/sheet1.xml':
                    content=content.replace(b'</sheetData>',b'<row r="4"><c r="B4"><f>_xlfn.DISPIMG("ID_PIC",1)</f></c></row></sheetData>')
                z.writestr(name,content)
            z.writestr('xl/cellimages.xml','<root xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><pic><cNvPr name="ID_PIC"/><a:blip r:embed="rId1"/></pic></root>')
            z.writestr('xl/_rels/cellimages.xml.rels','<Relationships><Relationship Id="rId1" Target="media/image1.png"/></Relationships>')
    pictures=script_media.xlsx_images(out.getvalue())
    assert len(pictures)==1 and set(pictures[0][1])=={'脚本!B4','脚本!D5'}
    assert parse_xlsx(out.getvalue()).uncovered_other==[]


@pytest.mark.parametrize('defect',['number','unknown_image','omitted_image','missing_rule'])
def test_invalid_multimodal_report_cannot_pass(client,monkeypatch,defect):
    task=setup(client); file_id=upload(client,workbook())
    def model(config,key,snapshot):
        value=report(snapshot)
        if defect=='number': value['findings'][1]['rule_number']=1
        if defect=='unknown_image': value['findings'][1]['image_ids']=['invented-image']
        if defect=='omitted_image': value['image_checks']=[]
        if defect=='missing_rule': value['findings'].pop()
        return {'report':value}
    monkeypatch.setattr(workflow,'call_script_model',model)
    review=submit(client,task,file_id).json()['data']['autoReview']
    assert review['status']=='failed' and not review['acceptance']['passed']


def test_image_only_workbook_review_and_uncovered_images_block_brand(client,monkeypatch):
    task=setup(client); file_id=upload(client,workbook(True))
    monkeypatch.setattr(workflow,'call_script_model',lambda c,k,s:{'report':report(s)})
    result=submit(client,task,file_id).json()['data']
    assert result['script']['extractionStatus']=='ready'
    assert result['autoReview']['status']=='succeeded' and result['autoReview']['acceptance']['passed']
    # A provider claiming compliance cannot erase local evidence of skipped media.
    monkeypatch.setattr(script_media,'MAX_IMAGES',0)
    other=upload(client,workbook())
    review=submit(client,task,other).json()['data']['autoReview']
    assert review['status']=='succeeded' and review['snapshot']['multimodal']['incomplete']
    assert review['acceptance']['status']=='incomplete' and not review['acceptance']['passed']


def test_model_sees_only_short_image_refs_not_64_char_hashes(client,monkeypatch):
    """回归：让模型逐字复述 64 位哈希会稳定出错（真实任务里一张图 id 第 33 位被改写，
    整次审核判 EVIDENCE_MISSING 失败）。模型只应看到 i1/i2 这类短引用。"""
    import httpx
    task=setup(client); file_id=upload(client,workbook())
    actual=httpx.Client; sent={}

    def provider(request):
        body=json.loads(request.content)
        sent['material']=json.loads(body['messages'][1]['content'][0]['text'])
        sent['labels']=[part['text'] for part in body['messages'][1]['content'] if part['type']=='text' and 'ref=' in part.get('text','')]
        value=report(sent['material'])
        return httpx.Response(200,json={'choices':[{'finish_reason':'stop','message':{'content':json.dumps(value,ensure_ascii=False)}}]})
    monkeypatch.setattr(workflow.httpx,'Client',lambda **kw:actual(transport=httpx.MockTransport(provider),**kw))
    review=submit(client,task,file_id).json()['data']['autoReview']
    assert review['status']=='succeeded',review
    ids={im['id'] for im in review['snapshot']['multimodal']['images']}
    dumped=json.dumps(sent['material'],ensure_ascii=False)
    assert not any(i in dumped for i in ids), '真实图片 id 不应出现在发给模型的材料里'
    assert [im['ref'] for im in sent['material']['multimodal']['images']] == ['i1','i2'][:len(sent['material']['multimodal']['images'])]
    assert all('ref=i' in label for label in sent['labels']), sent['labels']
    # 落库报告仍是真实 id，前端定位与图片预览不受影响
    assert {c['image_id'] for c in review['result']['report']['image_checks']} == ids
