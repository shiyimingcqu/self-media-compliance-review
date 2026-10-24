"""Read-only preview must preserve pages, ownership, source integrity and zero AI calls."""
import io

from docx import Document
from pypdf import PdfWriter

from app.core.database import connection
from app.modules.foundation.files_service import get_asset
from test_numbered_multimodal_script import setup
from test_script_versions import K


def upload_script(client, task, name, payload):
    response = client.post('/api/v1/files', files={'file': (name, payload)}, data={'purpose': 'script'}, headers={'Idempotency-Key': K()})
    assert response.status_code == 201, response.text
    asset = response.json()['data']
    response = client.post(f'/api/v1/tasks/{task}/scripts', json={'inputType': 'file', 'fileId': asset['id']}, headers={'Idempotency-Key': K()})
    assert response.status_code == 201, response.text
    script = response.json()['data']
    return asset, script, f"/api/v1/tasks/{task}/scripts/{script['scriptVersionId']}/layout"


def pdf_bytes():
    writer = PdfWriter();writer.add_blank_page(width=595, height=842);writer.add_blank_page(width=842, height=595)
    output = io.BytesIO();writer.write(output)
    return output.getvalue()


def test_docx_original_page_preview_preserves_geometry_without_changing_source(client, monkeypatch, tmp_path):
    monkeypatch.setenv('REVIEW_DOCX_PREVIEW','pdf')
    task = setup(client)
    document = Document();document.add_paragraph('产品介绍');document.add_table(rows=2, cols=3);document.add_page_break();document.add_paragraph('第二页')
    source = io.BytesIO();document.save(source)
    pdf = tmp_path / 'converted.pdf';pdf.write_bytes(pdf_bytes())
    calls = []
    def convert(payload, extension, target, cache_dir):
        calls.append((payload, extension, target));return pdf
    monkeypatch.setattr('app.modules.brief_script.office_preview.convert_office', convert)
    asset, script, url = upload_script(client, task, '脚本.docx', source.getvalue())
    response = client.get(url)
    assert response.status_code == 200, response.text
    original = response.json()['data']['original']
    assert original['kind'] == 'pages' and original['format'] == 'docx'
    assert original['pages'] == [{'page':1,'width':595,'height':842}, {'page':2,'width':842,'height':595}]
    for page in [1,2]:
        response = client.get(url + f'/pages/{page}')
        assert response.headers['content-type'] == 'image/png'
        assert response.content.startswith(b'\x89PNG')
    assert client.get(url+'/pages/0').status_code == 404
    assert client.get(url+'/pages/3').status_code == 404
    assert all(ext == 'docx' and target == 'pdf' for _, ext, target in calls)
    _, path = get_asset(client.app.state.settings, asset['id'])
    assert path.read_bytes() == source.getvalue()
    with connection(client.app.state.settings.database_path) as db:
        assert db.execute('SELECT count(*) FROM workflow_script_reviews').fetchone()[0] == 0
        project = db.execute('SELECT project_id FROM creator_tasks WHERE id=?',(task,)).fetchone()[0]
    other = client.post(f'/api/v1/projects/{project}/creators', json={'creatorName':'其他达人','platform':'douyin'}, headers={'Idempotency-Key':K()}).json()['data']['id']
    assert client.get(url.replace(task, other)).status_code == 404
    assert client.get((url+'/pages/1').replace(task, other)).status_code == 404
    path.write_bytes(b'changed')
    assert client.get(url).status_code == 409
    assert client.get(url+'/pages/1').status_code == 409


def test_pdf_original_can_be_read_even_if_scanned_text_is_unavailable(client):
    task = setup(client)
    _, script, url = upload_script(client, task, '扫描脚本.pdf', pdf_bytes())
    assert script['extractionStatus'] == 'invalid'  # Seeing the file is not text extraction success.
    assert script['parseMeta']['partial']
    assert len(client.get(url).json()['data']['original']['pages']) == 2
    assert client.get(url+'/pages/2').status_code == 200


def test_txt_preview_preserves_whitespace_and_utf16(client):
    task = setup(client)
    text = '  产品介绍\r\n\t镜头：咖啡  \r\n'
    _, _, url = upload_script(client, task, '脚本.txt', text.encode('utf-16'))
    assert client.get(url).json()['data']['original'] == {'kind':'text','format':'txt','text':text}


def test_converter_failure_is_visible_and_does_not_destroy_parsed_word(client, monkeypatch):
    monkeypatch.setenv('REVIEW_DOCX_PREVIEW','pdf')
    from app.core.errors import ApiError
    task = setup(client)
    document=Document();document.add_paragraph('产品介绍');source=io.BytesIO();document.save(source)
    _, script, url=upload_script(client,task,'脚本.docx',source.getvalue())
    def failed(*args, **kwargs):raise ApiError(422,'OFFICE_CONVERSION_FAILED','文件转换失败')
    monkeypatch.setattr('app.modules.brief_script.office_preview.convert_office',failed)
    assert client.get(url).json()['error']['code']=='OFFICE_CONVERSION_FAILED'
    assert script['extractionStatus']=='ready' and script['sourceMap'][0]['text']=='产品介绍'


def test_legacy_doc_registers_converted_text_but_retains_original(client, monkeypatch, tmp_path):
    task=setup(client)
    document=Document();document.add_paragraph('旧 Word 产品介绍');target=tmp_path/'converted.docx';document.save(target)
    monkeypatch.setattr('app.modules.brief_script.office_preview.convert_office',lambda *args,**kwargs:target)
    source=bytes.fromhex('d0cf11e0a1b11ae1')+b'legacy-fixture'
    asset,script,_=upload_script(client,task,'旧版.doc',source)
    assert script['extractionStatus']=='ready' and script['sourceMap'][0]['text']=='旧 Word 产品介绍'
    assert get_asset(client.app.state.settings,asset['id'])[1].read_bytes()==source


def test_docx_browser_original_is_available_without_conversion(client,monkeypatch):
    task=setup(client)
    document=Document();document.add_paragraph('原文件排版');document.add_table(rows=2,cols=2)
    source=io.BytesIO();document.save(source)
    def forbidden(*args,**kwargs):raise AssertionError('Browser preview must not launch Office')
    monkeypatch.setattr('app.modules.brief_script.office_preview.convert_office',forbidden)
    asset,_,url=upload_script(client,task,'原件.docx',source.getvalue())
    assert client.get(url).json()['data']['original']=={'kind':'docx','format':'docx'}
    response=client.get(url+'/document')
    assert response.headers['X-Preview-Source-Anchors']=='5'
    assert Document(io.BytesIO(response.content)).paragraphs[0].text=='原文件排版'
    assert get_asset(client.app.state.settings,asset['id'])[1].read_bytes()==source.getvalue()
    assert response.headers['content-type'].startswith('application/vnd.openxmlformats-officedocument.wordprocessingml.document')


def test_legacy_xls_uses_same_converted_sheet_for_parsing_and_layout(client,monkeypatch,tmp_path):
    from openpyxl import Workbook
    task=setup(client)
    workbook=Workbook();sheet=workbook.active;sheet.title='分镜';sheet.merge_cells('A1:B1');sheet['A1']='产品特写'
    target=tmp_path/'converted.xlsx';workbook.save(target);workbook.close()
    monkeypatch.setattr('app.modules.brief_script.office_preview.convert_office',lambda *args,**kwargs:target)
    _,script,url=upload_script(client,task,'旧表格.xls',bytes.fromhex('d0cf11e0a1b11ae1')+b'legacy-xls')
    assert script['extractionStatus']=='ready' and script['sourceMap'][0]['text']=='产品特写'
    data=client.get(url).json()['data']
    assert data['supported'] and data['sheets'][0]['name']=='分镜'
    assert data['sheets'][0]['merges']==[{'r1':1,'c1':1,'r2':1,'c2':2}]


def test_word_spacing_adaptation_is_preview_only_and_preserves_original(client):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from lxml import etree
    import zipfile
    task=setup(client)
    document=Document();paragraph=document.add_paragraph('第一行\n第二行')
    spacing=OxmlElement('w:spacing');spacing.set(qn('w:line'),'18pt');spacing.set(qn('w:lineRule'),'auto')
    paragraph._p.get_or_add_pPr().append(spacing)
    source=io.BytesIO();document.save(source)
    asset,script,url=upload_script(client,task,'带单位的行距.docx',source.getvalue())
    response=client.get(url+'/document')
    assert response.status_code==200 and response.headers['X-Preview-Formatting-Normalized']=='1'
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        root=etree.fromstring(archive.read('word/document.xml'))
        item=root.find('.//'+qn('w:spacing'))
        assert item.get(qn('w:line'))=='360' and item.get(qn('w:lineRule'))=='exact'
    assert script['sourceMap'][0]['text']=='第一行\n第二行'
    assert get_asset(client.app.state.settings,asset['id'])[1].read_bytes()==source.getvalue()


def test_source_pdf_pages_can_be_previewed_without_visual_ai_or_text_recognition(client):
    response=client.post('/api/v1/files',files={'file':('Brief.pdf',pdf_bytes())},data={'purpose':'brief'},headers={'Idempotency-Key':K()})
    asset=response.json()['data'];url=f"/api/v1/files/{asset['id']}/preview/pages"
    response=client.get(url+'/2')
    assert response.status_code==200 and response.content.startswith(b'\x89PNG')
    assert client.get(url+'/0').status_code==404 and client.get(url+'/3').status_code==404
    path=get_asset(client.app.state.settings,asset['id'])[1];path.write_bytes(b'changed')
    assert client.get(url+'/1').status_code==409
    text=client.post('/api/v1/files',files={'file':('Brief.txt',b'plain text')},data={'purpose':'brief'},headers={'Idempotency-Key':K()}).json()['data']
    assert client.get(f"/api/v1/files/{text['id']}/preview/pages/1").status_code==415
